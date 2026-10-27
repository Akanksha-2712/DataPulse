"""Unit tests for Stage 6: Report Generation and Multi-Format Export Module.

Tests cover:
- CSV stream generation for issues, anomalies, and drift.
- Microsoft Excel (.xlsx) generation via openpyxl:
  - Multi-tab presence (Summary, Quality Issues, Drift Results, Anomalies, Dataset Profile, Metadata)
  - Styling, auto-filters, freeze panes, number formats
  - Graceful degradation when components are omitted or None
- Professional PDF generation via ReportLab Platypus:
  - Valid %PDF- header byte signature
  - Multi-page document compilation with NumberedCanvas
  - Graceful degradation when optional components are None
- Filesystem export helper (save_report_to_disk)
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any, Dict

import numpy as np
import openpyxl
import pandas as pd
import pytest

from src.anomaly_detection import AnomalyDetectionResult, detect_anomalies
from src.drift_detection import DriftAnalysisResult, detect_drift
from src.profiling import DatasetProfile, profile_dataset
from src.quality_score import QualityScorecard, calculate_data_quality_score
from src.report_generator import (
    NumberedCanvas,
    ReportMetadata,
    export_analysis_to_excel,
    export_analysis_to_pdf,
    export_anomalies_to_csv,
    export_drift_to_csv,
    export_issues_to_csv,
    save_report_to_disk,
)
from src.synthetic_generator import generate_baseline_dataset, generate_current_dataset
from src.validation import ValidationIssue, ValidationResult, load_validation_rules, validate_dataset


@pytest.fixture(scope="module")
def sample_analysis_bundle() -> Dict[str, Any]:
    """Generate a complete end-to-end analysis bundle for report verification."""
    b_df = generate_baseline_dataset(n_rows=150, seed=42)
    c_df, _ = generate_current_dataset(n_rows=150, seed=101)

    # Load configuration
    config_path = Path("config/validation_rules.json")
    if not config_path.exists():
        # Fallback to parent if running in subfolder
        config_path = Path(__file__).resolve().parent.parent / "config" / "validation_rules.json"

    rules = load_validation_rules(config_path)

    prof = profile_dataset(c_df)
    val = validate_dataset(c_df, rules)
    dqs = calculate_data_quality_score(c_df, prof, val)
    drift = detect_drift(b_df, c_df)
    anom = detect_anomalies(c_df, b_df, contamination=0.08, random_state=42)
    meta = ReportMetadata(
        dataset_name="demo_current_sample.csv",
        reference_name="demo_baseline_sample.csv",
        worksheet_name="ProductionRecords",
        author="Quality Assurance Test Suite",
    )

    return {
        "profile": prof,
        "quality_scorecard": dqs,
        "validation_result": val,
        "drift_result": drift,
        "anomaly_result": anom,
        "metadata": meta,
    }


# ==============================================================================
# CSV EXPORT TESTS
# ==============================================================================

def test_export_issues_to_csv_with_data(sample_analysis_bundle: Dict[str, Any]) -> None:
    """Verify CSV export for validation issues produces well-formed CSV with expected header."""
    val: ValidationResult = sample_analysis_bundle["validation_result"]
    csv_bytes = export_issues_to_csv(val)

    assert isinstance(csv_bytes, bytes)
    assert len(csv_bytes) > 0

    content = csv_bytes.decode("utf-8")
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)

    # Check header
    expected_header = [
        "issue_index",
        "severity",
        "issue_type",
        "column",
        "row_reference",
        "observed_value",
        "description",
    ]
    assert rows[0] == expected_header
    assert len(rows) == val.total_issues + 1


def test_export_issues_to_csv_none_or_empty() -> None:
    """Verify CSV export gracefully handles None validation results without error."""
    csv_bytes = export_issues_to_csv(None)
    content = csv_bytes.decode("utf-8")
    rows = list(csv.reader(io.StringIO(content)))

    assert len(rows) == 1
    assert rows[0][0] == "issue_index"


def test_export_anomalies_to_csv_with_data(sample_analysis_bundle: Dict[str, Any]) -> None:
    """Verify CSV export for multivariate anomalies produces well-formed records."""
    anom: AnomalyDetectionResult = sample_analysis_bundle["anomaly_result"]
    csv_bytes = export_anomalies_to_csv(anom)

    content = csv_bytes.decode("utf-8")
    rows = list(csv.reader(io.StringIO(content)))

    assert rows[0] == [
        "row_index",
        "is_anomaly",
        "suspicion_score",
        "raw_decision_score",
        "top_deviating_features",
        "analyzed_values",
    ]
    assert len(rows) == len(anom.anomaly_records) + 1


def test_export_drift_to_csv_with_data(sample_analysis_bundle: Dict[str, Any]) -> None:
    """Verify CSV export for drift produces statistical test details."""
    drift: DriftAnalysisResult = sample_analysis_bundle["drift_result"]
    csv_bytes = export_drift_to_csv(drift)

    content = csv_bytes.decode("utf-8")
    rows = list(csv.reader(io.StringIO(content)))

    assert rows[0] == [
        "column_name",
        "column_type",
        "method",
        "engine",
        "statistic",
        "p_value",
        "threshold_alpha",
        "drift_detected",
        "sample_size_reference",
        "sample_size_current",
        "skipped",
        "skip_reason",
    ]
    assert len(rows) == len(drift.column_drift) + 1


# ==============================================================================
# EXCEL WORKBOOK EXPORT TESTS
# ==============================================================================

def test_export_analysis_to_excel_full_bundle(sample_analysis_bundle: Dict[str, Any]) -> None:
    """Verify openpyxl workbook creation, sheet structure, and cell formatting."""
    b = sample_analysis_bundle
    excel_bytes = export_analysis_to_excel(
        profile=b["profile"],
        quality_scorecard=b["quality_scorecard"],
        validation_result=b["validation_result"],
        drift_result=b["drift_result"],
        anomaly_result=b["anomaly_result"],
        metadata=b["metadata"],
    )

    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 2000

    # Reload with openpyxl to verify valid XLSX syntax and worksheet contents
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames

    expected_sheets = ["Summary", "Quality Issues", "Drift Results", "Anomalies", "Dataset Profile", "Metadata"]
    for s_name in expected_sheets:
        assert s_name in sheet_names, f"Missing expected sheet: {s_name}"

    # Verify Summary Sheet contents
    ws_sum = wb["Summary"]
    assert "DATA QUALITY MONITORING SYSTEM" in str(ws_sum["A1"].value)
    assert ws_sum.views.sheetView[0].showGridLines is True

    # Verify Quality Issues Sheet
    ws_issues = wb["Quality Issues"]
    assert ws_issues["A1"].value == "Issue ID"
    assert ws_issues["B1"].value == "Severity"
    assert ws_issues.freeze_panes == "A2"
    assert ws_issues.auto_filter.ref is not None
    assert ws_issues.max_row > 1

    # Verify Drift Results Sheet
    ws_drift = wb["Drift Results"]
    assert ws_drift["A1"].value == "Column Name"
    assert ws_drift.freeze_panes == "A2"
    assert ws_drift.max_row > 1

    # Verify Anomalies Sheet
    ws_anom = wb["Anomalies"]
    assert ws_anom["A1"].value == "Row Index"
    assert ws_anom.freeze_panes == "A2"
    assert ws_anom.max_row > 1

    # Verify Dataset Profile Sheet
    ws_prof = wb["Dataset Profile"]
    assert ws_prof["A1"].value == "Column Name"
    assert ws_prof.max_row > 1

    # Verify Metadata Sheet
    ws_meta = wb["Metadata"]
    assert ws_meta["A1"].value == "TECHNICAL AUDIT TRAIL & ACADEMIC METHODOLOGY"


def test_export_analysis_to_excel_graceful_degradation() -> None:
    """Verify openpyxl workbook exports successfully even when all optional inputs are None."""
    excel_bytes = export_analysis_to_excel(
        profile=None,
        quality_scorecard=None,
        validation_result=None,
        drift_result=None,
        anomaly_result=None,
        metadata=None,
    )

    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert len(wb.sheetnames) == 6
    assert "Summary" in wb.sheetnames
    ws_issues = wb["Quality Issues"]
    assert ws_issues.max_row >= 2


# ==============================================================================
# PDF REPORT EXPORT TESTS
# ==============================================================================

def test_export_analysis_to_pdf_full_bundle(sample_analysis_bundle: Dict[str, Any]) -> None:
    """Verify ReportLab Platypus PDF compilation, signature, and multi-page flow."""
    b = sample_analysis_bundle
    pdf_bytes = export_analysis_to_pdf(
        profile=b["profile"],
        quality_scorecard=b["quality_scorecard"],
        validation_result=b["validation_result"],
        drift_result=b["drift_result"],
        anomaly_result=b["anomaly_result"],
        metadata=b["metadata"],
    )

    assert isinstance(pdf_bytes, bytes)
    # Check valid PDF header
    assert pdf_bytes.startswith(b"%PDF-")
    # Multi-page academic report with all sections should exceed 8 KB
    assert len(pdf_bytes) > 8000


def test_export_analysis_to_pdf_graceful_degradation() -> None:
    """Verify PDF compiles cleanly when no reference data, rules, or anomalies are provided."""
    pdf_bytes = export_analysis_to_pdf(
        profile=None,
        quality_scorecard=None,
        validation_result=None,
        drift_result=None,
        anomaly_result=None,
        metadata=None,
    )

    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF-")
    assert len(pdf_bytes) > 2000


# ==============================================================================
# DISK PERSISTENCE & METADATA TESTS
# ==============================================================================

def test_save_report_to_disk(tmp_path: Path) -> None:
    """Verify save_report_to_disk creates parent directories and writes identical bytes."""
    dest = tmp_path / "nested" / "reports" / "test_report.pdf"
    dummy_bytes = b"%PDF-1.4 Mock PDF payload bytes for testing"

    saved_path = save_report_to_disk(dummy_bytes, dest)
    assert saved_path.exists()
    assert saved_path.is_file()
    assert saved_path.read_bytes() == dummy_bytes


def test_report_metadata_to_dict() -> None:
    """Verify ReportMetadata serialization to dictionary."""
    meta = ReportMetadata(
        dataset_name="sales_2026.csv",
        reference_name="sales_2025.csv",
        worksheet_name="Q1",
        author="Data Scientist",
    )
    d = meta.to_dict()
    assert d["dataset_name"] == "sales_2026.csv"
    assert d["reference_name"] == "sales_2025.csv"
    assert d["worksheet_name"] == "Q1"
    assert d["author"] == "Data Scientist"
    assert isinstance(d["generated_at"], str)
